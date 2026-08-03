from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from karverktyg.config.models import TransitionKind
from karverktyg.export import ChangelistAckRequired, build_changelist, reconcile
from karverktyg.scoutnet.models import Member, MemberList
from karverktyg.uppflyttning.models import MasterSet, MoveEntry, MoveStatus

WHEN = datetime(2026, 8, 3, 9, 0, 0)


def _entry(no, name, tgt, tgtid, status=MoveStatus.READY):
    return MoveEntry(
        member_no=no, member_name=name, birth_year=2016,
        source_avdelning="Hajarna", source_troop_id=10155,
        target_avdelning=tgt, target_troop_id=tgtid,
        transition=TransitionKind.SAME_WEEKDAY, status=status,
    )


def _master(entries):
    return MasterSet(cohort_year=2026, entries=entries)


def test_changelist_requires_ack_for_off_cohort():
    master = _master([
        _entry("1", "Alva Ek", "Kämparna", 10158),
        _entry("2", "Öjvind Ö", None, None, status=MoveStatus.OFF_COHORT),
    ])
    with pytest.raises(ChangelistAckRequired):
        build_changelist(master, kar_name="Finn", term_label="Höst 2026",
                         generated_at=WHEN, config_version="v1")
    # with an acknowledgement it succeeds
    data = build_changelist(master, kar_name="Finn", term_label="Höst 2026",
                            generated_at=WHEN, config_version="v1",
                            ack_by="alex", ack_at=WHEN)
    assert isinstance(data, bytes) and data[:2] == b"PK"  # xlsx is a zip


def test_changelist_structure_and_ordering():
    master = _master([
        _entry("300", "Örjan Ö", "Kämparna", 10158),
        _entry("100", "Alva Ek", "Kämparna", 10158),
        _entry("200", "Bo Berg", "Spejarna", 10161),
    ])
    data = build_changelist(master, kar_name="Finn", term_label="Höst 2026",
                            generated_at=WHEN, config_version="v1")
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames[0] == "Översikt"
    assert "Kämparna" in wb.sheetnames and "Spejarna" in wb.sheetnames

    ws = wb["Kämparna"]
    assert ws["A1"].value == "Medlemsnummer"  # member_no first column
    assert ws.freeze_panes == "A2"
    # Swedish collation within the sheet: Alva (A) before Örjan (Ö)
    assert ws["A2"].value == "100"
    assert ws["A3"].value == "300"


def test_reconcile_intent_based():
    master = _master([
        _entry("m1", "A", "Kämparna", 200),
        _entry("m2", "B", "Kämparna", 200),
        _entry("m3", "C", "Spejarna", 300),
    ])
    after = MemberList(members=[
        Member(member_no="m1", unit="Kämparna", unit_troop_id=200),   # applied
        Member(member_no="m3", unit="Rockorna", unit_troop_id=999),   # elsewhere
        # m2 missing -> not found
        Member(member_no="resident", unit="Kämparna", unit_troop_id=200),  # not drift
    ])
    res = reconcile(master, after)
    assert res.summary() == {"applied": 1, "not_found": 1, "elsewhere": 1}
    assert res.applied[0].member_no == "m1"
    assert res.not_found[0].member_no == "m2"
    assert res.elsewhere[0].member_no == "m3"
    assert not res.all_applied

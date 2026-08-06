from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from scoutnet_membership_manager.export import build_affected_scouts_xlsx
from scoutnet_membership_manager.export.affected_scouts_xlsx import _HEADERS
from scoutnet_membership_manager.scoutnet.models import Member, MemberList, Role
from scoutnet_membership_manager.settings import Mode, Settings
from scoutnet_membership_manager.uppflyttning.models import MasterSet, MoveEntry, MoveStatus
from scoutnet_membership_manager.web import create_app


def _m(no, first, last, unit, **kw):
    return Member(
        member_no=no,
        first_name=first,
        last_name=last,
        unit=unit,
        emails=kw.get("emails", {}),
        phones=kw.get("phones", {}),
        guardian_names=kw.get("guardian_names", {}),
    )


def _entry(no, name="X X"):
    return MoveEntry(
        member_no=no,
        member_name=name,
        birth_year=2016,
        source_avdelning="Upptäckarna",
        source_troop_id=10155,
        target_avdelning="Äventyrarna",
        target_troop_id=10158,
        transition="merge",
        status=MoveStatus.READY,
    )


def _sheet(master, memberlist):
    wb = load_workbook(BytesIO(build_affected_scouts_xlsx(master, memberlist)))
    # The contact roster, plus a follow-up sheet for patrol roles to end by hand.
    assert wb.sheetnames == ["Berörda scouter", "Patrulledare att avsluta"]
    return wb["Berörda scouter"]


def test_headers_are_contacts_only_no_ssn_or_address_and_neutral():
    ws = _sheet(MasterSet(cohort_year=2026), MemberList(members=[]))
    headers = [c.value for c in ws[1]]
    assert headers == _HEADERS
    joined = " ".join(headers).lower()
    assert "personnummer" not in joined and "adress" not in joined
    # Neutral guardian labels — no mor/far asserted.
    assert "mor" not in joined and "far" not in joined
    assert "Avdelning" in headers
    assert "Medlemmens e-post" in headers and "Medlemmens mobil" in headers
    assert "Anhörig 1 – e-post" in headers and "Anhörig 2 – telefon" in headers


def test_contacts_join_and_current_avdelning():
    m = _m(
        "1",
        "Alva",
        "Ek",
        "Upptäckarna",
        emails={"contact_email": "alva@example.org", "contact_email_mum": "mor@example.org"},
        phones={"contact_mobile_phone": "072", "contact_mobile_mum": "070"},
        guardian_names={"contact_mothers_name": "Mor Ek"},
    )
    ml = MemberList(members=[m])
    master = MasterSet(cohort_year=2026, entries=[_entry("1", "Alva Ek")])
    ws = _sheet(master, ml)
    row = [c.value for c in ws[2]]
    assert row[0] == "1"
    assert row[1] == "Upptäckarna"  # current avdelning, not the target
    assert row[2] == "Alva Ek"
    # member own: e-post, mobil ; guardian 1: namn, e-post, telefon
    assert row[3] == "alva@example.org" and row[4] == "072"
    assert row[5] == "Mor Ek" and row[6] == "mor@example.org" and row[7] == "070"


def test_sole_guardian_in_dad_field_lands_in_first_column():
    # Only the "dad" fields are set — the sole guardian must fill Anhörig 1, not 2.
    m = _m(
        "1",
        "Alva",
        "Ek",
        "Upptäckarna",
        emails={"contact_email_dad": "far@example.org"},
        phones={"contact_mobile_dad": "070"},
        guardian_names={"contact_fathers_name": "Far Ek"},
    )
    ml = MemberList(members=[m])
    master = MasterSet(cohort_year=2026, entries=[_entry("1", "Alva Ek")])
    row = [c.value for c in _sheet(master, ml)[2]]
    assert row[5] == "Far Ek" and row[6] == "far@example.org" and row[7] == "070"
    # Anhörig 2 stays empty.
    assert row[8] in (None, "") and row[9] in (None, "") and row[10] in (None, "")


def test_telephone_falls_back_when_no_mobile():
    m = _m(
        "1",
        "Alva",
        "Ek",
        "Upptäckarna",
        phones={"contact_telephone_mum": "08-123"},
        guardian_names={"contact_mothers_name": "Mor Ek"},
    )
    ml = MemberList(members=[m])
    master = MasterSet(cohort_year=2026, entries=[_entry("1", "Alva Ek")])
    row = [c.value for c in _sheet(master, ml)[2]]
    assert row[5] == "Mor Ek" and row[7] == "08-123"


def test_orders_by_avdelning_then_surname_and_skips_unknown_members():
    ml = MemberList(
        members=[
            _m("1", "Berit", "Berg", "Upptäckarna"),
            _m("2", "Alva", "Alm", "Upptäckarna"),
            _m("3", "Cesar", "Ek", "Spårarna"),
        ]
    )
    master = MasterSet(
        cohort_year=2026,
        # entry "99" has no matching member and must be skipped, not error.
        entries=[_entry("1"), _entry("2"), _entry("3"), _entry("99")],
    )
    ws = _sheet(master, ml)
    names = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
    # Spårarna before Upptäckarna; within Upptäckarna, Alm before Berg.
    assert names == ["Cesar Ek", "Alva Alm", "Berit Berg"]


def test_patrol_roles_sheet_lists_holders_with_scoutnet_links():
    m = _m("1", "Alva", "Ek", "Utmanarna")
    m.patrol = "Örnen"
    m.roles = [
        Role(scope="patrol", scope_id=500, role_id=2, role_key="leader", role_name="Patrulledare"),
        # A troop-scoped adult role must NOT appear — only patrol-scoped roles.
        Role(scope="troop", scope_id=10, role_id=9, role_key="leader", role_name="Ledare"),
    ]
    ml = MemberList(members=[m])
    master = MasterSet(cohort_year=2026, entries=[_entry("1", "Alva Ek")])
    wb = load_workbook(BytesIO(build_affected_scouts_xlsx(master, ml)))
    ws = wb["Patrulledare att avsluta"]
    assert [c.value for c in ws[1]] == ["Namn", "Avdelning", "Patrull", "Uppdrag", "Scoutnet-länk"]
    row = [c.value for c in ws[2]]
    assert row[0] == "Alva Ek" and row[1] == "Utmanarna" and row[2] == "Örnen"
    assert row[3] == "Patrulledare"
    link = ws.cell(row=2, column=5)
    assert link.value == "https://scoutnet.se/organisation/user/1"
    assert link.hyperlink.target == "https://scoutnet.se/organisation/user/1"
    assert ws.max_row == 2  # the troop-scoped Ledare role is excluded


def test_api_affected_scouts_xlsx():
    settings = Settings(mode=Mode.FIXTURE, database_url="sqlite://", cohort_year=None)
    r = create_app(settings).test_client().get("/api/uppflyttning/berorda-scouter.xlsx?group=all")
    assert r.status_code == 200 and r.data[:2] == b"PK"  # xlsx is a zip

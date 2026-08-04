from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from karverktyg.export.dues_xlsx import build_dues_xlsx
from karverktyg.scoutnet.models import DueDate, Member, MemberList
from karverktyg.scoutnet.parse import parse_due_date
from karverktyg.settings import Mode, Settings
from karverktyg.web import create_app


def _m(no, first, last, unit, prev_code, cur_code="not_invoiced", **kw):
    return Member(
        member_no=no,
        first_name=first,
        last_name=last,
        unit=unit,
        prev_term_code=prev_code,
        current_term_code=cur_code,
        prev_term_label="Vår 2026",
        current_term_label="Höst 2026",
        prev_term_value=kw.get("prev_value", ""),
        current_term_value=kw.get("cur_value", ""),
        prev_term_due=kw.get("prev_due", DueDate()),
        kid=kw.get("kid"),
        emails=kw.get("emails", {}),
        phones=kw.get("phones", {}),
        guardian_names=kw.get("guardian_names", {}),
    )


def _list(members):
    return MemberList(members=members, prev_term_label="Vår 2026", current_term_label="Höst 2026")


def _sheet(members):
    wb = load_workbook(BytesIO(build_dues_xlsx(_list(members))))
    assert wb.sheetnames == ["Medlemsavgifter"]  # a single flat sheet with everyone
    return wb["Medlemsavgifter"]


def _values(ws):
    return [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]


# --- Due-date parsing (§4 / §16 DoD) --------------------------------------


def test_parse_due_date_retains_both_dates():
    shifted = parse_due_date("2026-04-30 (2026-02-28)")
    assert shifted.current == "2026-04-30" and shifted.original == "2026-02-28"
    assert shifted.shifted and shifted.display() == "2026-04-30 (2026-02-28)"
    plain = parse_due_date("2026-02-28")
    assert plain.current == "2026-02-28" and plain.original is None and not plain.shifted
    empty = parse_due_date("")
    assert empty.current is None and empty.display() == ""


# --- Bucket routing (§4 / §19) --------------------------------------------


def test_not_invoiced_and_paid_are_excluded():
    ws = _sheet(
        [_m("1", "A", "A", "Hajarna", "not_invoiced"), _m("2", "B", "B", "Hajarna", "paid")]
    )
    assert ws.max_row == 1  # header only — nobody to chase


def test_outstanding_included_in_the_single_sheet():
    ws = _sheet([_m("1", "A", "A", "Vikingarna", "unpaid_overdue_reminded")])
    assert ws.max_row == 2 and ws["A2"].value == "1" and ws["D2"].value == "Vikingarna"


def test_unknown_code_included_but_flagged():
    ws = _sheet(
        [
            _m("1", "A", "A", "Hajarna", "some_future_code"),  # unrecognised
            _m("2", "B", "B", "Hajarna", "unpaid_overdue_reminded"),  # normal chase
        ]
    )
    ids = _values(ws)
    assert "1" in ids and "2" in ids  # both in the one sheet now
    assert any(isinstance(v, str) and "Okänd betalkod" in v for v in ids)  # unknown flagged


def test_partial_credit_included_and_distinguished():
    ws = _sheet(
        [
            _m(
                "9",
                "A",
                "A",
                "Hajarna",
                "paid_partial_credit",
                prev_value="Betalt, delvis krediterat",
            )
        ]
    )
    values = _values(ws)
    assert "paid_partial_credit" in values
    assert any(isinstance(v, str) and "Felaktigt belopp" in v for v in values)


# --- Layout: one sheet, all avdelningar, sorted (§19) ----------------------


def test_all_avdelningar_in_one_sheet_sorted():
    ws = _sheet(
        [
            _m("1", "Bo", "Berg", "Rockorna", "unpaid_overdue_reminded"),
            _m("2", "B", "B", "Ledare", "unpaid_overdue_reminded"),  # adults owe dues too
            _m("3", "C", "C", None, "unpaid_overdue_reminded"),  # no avdelning -> last
            _m("4", "Alva", "Alm", "Rockorna", "unpaid_overdue_reminded"),
        ]
    )
    avdelningar = [ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)]
    # Grouped by avdelning (Ledare, Rockorna, then the no-avdelning bucket last).
    assert avdelningar == ["Ledare", "Rockorna", "Rockorna", "(ingen avdelning)"]
    # Within Rockorna, surname order (Alm before Berg).
    assert ws["B3"].value == "Alm" and ws["B4"].value == "Berg"


def test_contact_and_due_columns_populate():
    m = _m(
        "1",
        "A",
        "A",
        "Hajarna",
        "unpaid_overdue_reminded",
        prev_due=DueDate(current="2026-04-30", original="2026-02-28"),
        kid="1234567",
        emails={"contact_email_mum": "mor@example.org", "contact_email": "egen@example.org"},
        phones={"contact_mobile_dad": "070", "contact_mobile_phone": "072"},
        guardian_names={"contact_mothers_name": "Mor Ek", "contact_fathers_name": "Far Ek"},
    )
    values = _values(_sheet([m]))
    assert "2026-04-30 (2026-02-28)" in values and "1234567" in values
    assert "mor@example.org" in values and "Mor Ek" in values and "Far Ek" in values
    assert "egen@example.org" in values


# --- Against the real fixture + API ---------------------------------------


def test_fixture_scale(memberlist):
    wb = load_workbook(BytesIO(build_dues_xlsx(memberlist)))
    # The fabricated demo has 19 members with an outstanding previous term, plus the header.
    assert wb["Medlemsavgifter"].max_row == 20


def test_api_dues_xlsx():
    settings = Settings(mode=Mode.FIXTURE, database_url="sqlite://", cohort_year=None)
    r = create_app(settings).test_client().get("/api/dues.xlsx")
    assert r.status_code == 200 and r.data[:2] == b"PK"

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from karverktyg.export.oversikt_xlsx import build_oversikt_xlsx
from karverktyg.oversikt import OversiktInputs, build_oversikt
from karverktyg.scoutnet.models import Member, MemberList, Role
from karverktyg.settings import Mode, Settings
from karverktyg.web import create_app

WHEN = datetime(2026, 8, 3, 9, 0, 0)


def _member(no, unit, utc, by=None, troop_id=None, roles=()):
    return Member(
        member_no=no,
        unit=unit,
        unit_troop_id=troop_id,
        unit_type_code=utc,
        birth_year=by,
        roles=list(roles),
    )


def _ml(members):
    return MemberList(members=members, current_term_label="Höst 2026", prev_term_label="Vår 2026")


def _leader_role(troop_id):
    return Role(scope="troop", scope_id=troop_id, role_id=1, role_key="leader", role_name="Ledare")


def _build(config, members, **kw):
    return build_oversikt(
        OversiktInputs(memberlist=_ml(members), config=config, cohort_year_n=None, **kw)
    )


def _find_avd(d, name):
    for g in d["composition"]["groups"]:
        for a in g["avdelningar"]:
            if a["name"] == name:
                return a
    return None


# --- Composition & registry (§20) -----------------------------------------


def test_empty_configured_avdelning_is_a_dash_row(config):
    # Only Hajarna has members; every other configured avdelning is empty.
    d = _build(config, [_member("1", "Hajarna", 2, by=2018, troop_id=10155)])
    viking = _find_avd(d, "Vikingarna")  # configured, no members
    assert viking is not None and viking["members"] == "-"  # not omitted, not 0
    hajarna = _find_avd(d, "Hajarna")
    assert hajarna["members"] == 1
    assert d["composition"]["kar_total"] == 1


def test_discovered_but_unconfigured_is_flagged(config):
    d = _build(
        config,
        [
            _member("1", "Hajarna", 2, by=2018, troop_id=10155),
            _member("2", "Nyavdelning", 2, by=2018, troop_id=98765),
        ],
    )
    nya = _find_avd(d, "Nyavdelning")
    assert nya is not None and nya["flag"] and "saknas i config" in nya["flag"]


def test_no_unit_members_counted_separately(config):
    d = _build(config, [_member("1", None, None), _member("2", "Hajarna", 2, by=2018)])
    assert d["composition"]["no_unit_count"] == 1
    assert d["composition"]["kar_total"] == 2


# --- Leaders & scouts-per-leader (§20) -------------------------------------


def test_two_leader_numbers_and_ratio_split(config):
    members = [
        _member("s1", "Hajarna", 2, by=2018, troop_id=10155),
        _member("s2", "Hajarna", 2, by=2018, troop_id=10155),
        _member("s3", "Hajarna", 2, by=2018, troop_id=10155),
        # an adult leader of Hajarna (own unit = Ledare)
        _member("a", "Ledare", 7, troop_id=10172, roles=[_leader_role(10155)]),
        # a youth leader of Hajarna (own unit = a scout avdelning)
        _member("y", "Kämparna", 3, by=2016, troop_id=10158, roles=[_leader_role(10155)]),
    ]
    d = _build(config, members)
    assert d["leaders"]["ledare_members"] == 1  # only the adult sits in Ledare
    assert d["leaders"]["role_holders"] == 2  # two hold a leader role
    haj = next(s for s in d["scouts_per_leader"] if s["avdelning"] == "Hajarna")
    assert haj["scouts"] == 3 and haj["leaders"] == 2
    assert haj["adult_leaders"] == 1 and haj["youth_leaders"] == 1
    assert haj["ratio"] == "1.5:1"


def test_scouts_per_leader_dash_when_no_leaders(config):
    d = _build(config, [_member("1", "Rockorna", 2, by=2018, troop_id=10157)])
    row = next(s for s in d["scouts_per_leader"] if s["avdelning"] == "Rockorna")
    assert row["leaders"] == "-" and row["ratio"] == "-"  # 0 leaders -> not applicable


# --- Projection (§20) ------------------------------------------------------


def test_projection_transitions_and_spring_autumn(config):
    # Oldest Spårare (born 2016 -> age 10 in 2026) move up; younger stay.
    movers = [_member(f"m{i}", "Hajarna", 2, by=2016, troop_id=10155) for i in range(3)]
    stayers = [_member(f"s{i}", "Hajarna", 2, by=2018, troop_id=10155) for i in range(2)]
    d = _build(config, movers + stayers)
    assert d["shift_applied"] is False and d["projection_targets"] == 2026
    sp_to_up = next(t for t in d["projection"]["transitions"] if t["from"] == "Spårare")
    assert sp_to_up["count"] == 3


def test_projection_autumn_targets_next_year(config):
    # Everyone correctly placed for their bracket -> no movers -> shift applied.
    d = _build(config, [_member(f"s{i}", "Hajarna", 2, by=2018, troop_id=10155) for i in range(3)])
    assert d["shift_applied"] is True and d["projection_targets"] == 2027


def test_sparare_recruitment_shows_all_three_and_negative_in_words(config):
    movers = [_member(f"m{i}", "Hajarna", 2, by=2016, troop_id=10155) for i in range(2)]
    # Four Spårare-aged applicants (born 2018) waiting -> Y=4 > X=2 -> Z negative.
    applicants = _ml([_member(f"w{i}", None, None, by=2018) for i in range(4)])
    d = _build(config, movers, applicants={"waiting": applicants, "awaiting_approval": None})
    sp = d["projection"]["spararrekrytering"]
    assert sp["x_leaving"] == 2 and sp["y_pending"] == 4 and sp["z_target"] == -2
    assert "inget rekryteringsbehov" in sp["sentence"]  # negative stated in words
    assert sp["provisional"] is True  # awaiting_approval missing


def test_utmanare_projects_static(config):
    d = _build(config, [_member("u", "Finndus", 5, by=2010, troop_id=20001)])
    row = next(r for r in d["projection"]["rows"] if r["avdelning"] == "Finndus")
    assert row["basis"] == "static" and row["next"] == row["current"]


# --- Reconciliation & KPIs (§20) ------------------------------------------


def test_reconciliation_disagreement_is_surfaced(config):
    d = _build(
        config, [_member("1", "Hajarna", 2, by=2018, troop_id=10155)], org_group={"membercount": 99}
    )
    assert d["reconciliation"]["membercount"]["agree"] is False


def test_share_paid_and_retention(config):
    members = [
        _member("1", "Hajarna", 2, by=2018, troop_id=10155),
        _member("2", "Hajarna", 2, by=2018, troop_id=10155),
    ]
    members[0].prev_term_code = "paid"
    d = _build(config, members)
    assert d["kpis"]["share_paid"]["computed_paid"] == 1
    assert d["kpis"]["retention"]["available"] is False  # unavailable, stated (§20)


# --- Export (§19) ----------------------------------------------------------


def test_xlsx_is_aggregate_only_with_dash(config):
    d = _build(config, [_member("HEMLIG123", "Hajarna", 2, by=2018, troop_id=10155)])
    wb = load_workbook(BytesIO(build_oversikt_xlsx(d, kar_name="Finn", generated_at=WHEN)))
    assert wb.sheetnames == ["Info", "Sammansättning", "Ledare", "Prognos", "Nyckeltal"]
    all_cells = [
        str(c.value)
        for name in wb.sheetnames
        for row in wb[name].iter_rows()
        for c in row
        if c.value is not None
    ]
    assert "HEMLIG123" not in all_cells  # aggregate-only: no member numbers
    assert "-" in all_cells  # empty avdelningar render as the literal dash


# --- API -------------------------------------------------------------------


def _client():
    return create_app(
        Settings(mode=Mode.FIXTURE, database_url="sqlite://", cohort_year=None)
    ).test_client()


def test_api_overview_rich_and_consistent():
    d = _client().get("/api/overview").get_json()
    assert d["member_count"] == 371  # legacy key preserved
    assert d["composition"]["kar_total"] == 371
    assert d["leaders"]["ledare_members"] == 120
    assert d["reconciliation"]["membercount"]["agree"] is True  # 371 == synthesised org count


def test_api_oversikt_exports():
    c = _client()
    assert c.get("/api/oversikt.xlsx").status_code == 200
    assert c.get("/api/oversikt.pdf").status_code in (
        200,
        503,
    )  # 200 in container, 503 without pdf extra

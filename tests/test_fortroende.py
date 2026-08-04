from __future__ import annotations

import importlib.util
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from scoutnet_membership_manager.config.models import ExpectedPost
from scoutnet_membership_manager.export.fortroende_xlsx import build_fortroende_xlsx
from scoutnet_membership_manager.export.report import PdfUnavailable, html_to_pdf, render_html
from scoutnet_membership_manager.fortroende import (
    Assignment,
    fortroendeuppdrag,
    group_by_section,
    rolecount_reconciliation,
)
from scoutnet_membership_manager.scoutnet.models import Member, MemberList, Role
from scoutnet_membership_manager.settings import Mode, Settings
from scoutnet_membership_manager.web import create_app

WHEN = datetime(2026, 8, 3, 9, 0, 0)
HAS_WEASYPRINT = importlib.util.find_spec("weasyprint") is not None


def _role(scope, role_key, role_name, scope_id=1025):
    return Role(scope=scope, scope_id=scope_id, role_id=1, role_key=role_key, role_name=role_name)


def _member(no, first, last, roles):
    return Member(member_no=no, first_name=first, last_name=last, roles=roles)


# --- Pure domain ----------------------------------------------------------


def test_only_group_scope_is_listed(config):
    ml = MemberList(
        members=[
            _member("1", "A", "A", [_role("group", "board_member", "Kårstyrelseledamot")]),
            _member("2", "B", "B", [_role("troop", "leader", "Avdelningsledare", scope_id=10155)]),
            _member("3", "C", "C", [_role("patrol", "leader", "Patrulledare", scope_id=42)]),
        ]
    )
    result = fortroendeuppdrag(ml, config)
    assert [a.role_key for a in result.assignments] == ["board_member"]  # only group
    assert result.group_count == 1
    assert result.total_roles_parsed == 3  # all scopes counted for the cross-check


def test_unknown_role_key_passes_through_sorted_last(config, memberlist):
    # The demo fixture carries role_key "key_responsible" -> "Nyckelansvarig", an
    # unfamiliar key: §18 pass-through must show it, never drop it, sorted last.
    result = fortroendeuppdrag(memberlist, config)
    keys = [a.role_key for a in result.assignments]
    names = [a.role_name for a in result.assignments]
    assert "Nyckelansvarig" in names
    # Kårordförande (group-scoped role_key "leader") sorts first (constitutional).
    assert result.assignments[0].role_name == "Kårordförande"
    # The unconfigured key sorts after every configured styrelse key.
    assert keys.index("key_responsible") > keys.index("board_member")


def test_styrelse_constitutional_order(config):
    ml = MemberList(
        members=[
            _member("1", "Z", "Z", [_role("group", "board_member", "Kårstyrelseledamot")]),
            _member("2", "Y", "Y", [_role("group", "treasurer", "Kårkassör")]),
            _member("3", "X", "X", [_role("group", "leader", "Kårordförande")]),
            _member("4", "W", "W", [_role("group", "hut_responsible", "Stugansvarig")]),
        ]
    )
    order = [a.role_key for a in fortroendeuppdrag(ml, config).assignments]
    assert order.index("leader") < order.index("treasurer") < order.index("board_member")
    assert order.index("board_member") < order.index("hut_responsible")  # unconfigured last


def test_name_collation_within_a_role(config):
    ml = MemberList(
        members=[
            _member("1", "Örjan", "Ö", [_role("group", "board_member", "Kårstyrelseledamot")]),
            _member("2", "Alva", "A", [_role("group", "board_member", "Kårstyrelseledamot")]),
            _member("3", "Bo", "B", [_role("group", "board_member", "Kårstyrelseledamot")]),
        ]
    )
    names = [a.member_name for a in fortroendeuppdrag(ml, config).assignments]
    assert names == ["Alva A", "Bo B", "Örjan Ö"]  # å/ä/ö sort after z (§2)


def test_label_override_applies(config):
    cfg = config.model_copy(update={"role_label_overrides": {"cottage_rental": "Stugvärd"}})
    ml = MemberList(
        members=[_member("1", "A", "A", [_role("group", "cottage_rental", "Stugbokare")])]
    )
    a = fortroendeuppdrag(ml, cfg).assignments[0]
    assert a.role_name == "Stugbokare" and a.label == "Stugvärd"


def test_empty_roles_contributes_nothing(config):
    ml = MemberList(members=[_member("1", "A", "A", [])])
    result = fortroendeuppdrag(ml, config)
    assert result.assignments == [] and result.total_roles_parsed == 0


def test_rolecount_reconciliation():
    assert rolecount_reconciliation(120, 120) == {
        "total_parsed": 120,
        "rolecount": 120,
        "available": True,
        "match": True,
    }
    assert rolecount_reconciliation(120, 118)["match"] is False
    unavailable = rolecount_reconciliation(120, None)
    assert unavailable["available"] is False and unavailable["match"] is False


def test_vacancies_are_opt_in(config):
    ml = MemberList(
        members=[_member("1", "A", "A", [_role("group", "board_member", "Kårstyrelseledamot")])]
    )
    # No expected posts configured -> no vacancies.
    assert fortroendeuppdrag(ml, config).vacancies == []
    # Expect 3 board members, 1 filled -> reported with missing=2.
    cfg = config.model_copy(
        update={"expected_fortroende": [ExpectedPost(role_key="board_member", count=3)]}
    )
    vac = fortroendeuppdrag(ml, cfg).vacancies
    assert len(vac) == 1 and vac[0].filled == 1 and vac[0].missing == 2


def test_three_sections_board_first_delegate_last(config):
    ml = MemberList(
        members=[
            _member("1", "A", "A", [_role("group", "district_voter", "Ombud distriktstämma")]),
            _member("2", "B", "B", [_role("group", "hut_responsible", "Stugansvarig")]),
            _member("3", "C", "C", [_role("group", "leader", "Kårordförande")]),
            # Utmanarscoutrepresentant is a board seat (§18), not "other".
            _member("4", "D", "D", [_role("group", "scout_challenge_rep", "Utmanarscoutrep")]),
        ]
    )
    result = fortroendeuppdrag(ml, config)
    assert [a.section for a in result.assignments] == ["board", "board", "other", "delegate"]
    board_keys = [a.role_key for a in result.assignments if a.section == "board"]
    assert "scout_challenge_rep" in board_keys  # board seat, not other/delegate
    labels = [s["label"] for s in group_by_section(result.assignments)]
    assert labels == ["Kårstyrelse", "Övriga förtroendeuppdrag", "Ombud och representanter"]


# --- Excel export ---------------------------------------------------------


def test_fortroende_xlsx_structure_and_order(config):
    result = fortroendeuppdrag(
        MemberList(
            members=[
                _member("2", "B", "B", [_role("group", "hut_responsible", "Stugansvarig")]),
                _member("1", "A", "A", [_role("group", "board_member", "Kårstyrelseledamot")]),
            ]
        ),
        config,
    )
    data = build_fortroende_xlsx(
        result, kar_name="Finn", term_label="Höst 2026", generated_at=WHEN, rolecount=5
    )
    assert data[:2] == b"PK"
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Info", "Förtroendeuppdrag"]
    ws = wb["Förtroendeuppdrag"]
    assert ws["A1"].value == "Sektion" and ws["B1"].value == "Roll"
    assert ws.freeze_panes == "A2"
    # board_member (Kårstyrelse) sorts before hut_responsible (Övriga); role_key in col E.
    assert ws["A2"].value == "Kårstyrelse"
    assert ws["E2"].value == "board_member" and ws["E3"].value == "hut_responsible"


# --- Report HTML / PDF guard ---------------------------------------------


def _sample_context():
    return {
        "kar": "Scoutkåren Finn",
        "term": "Höst 2026",
        "generated": WHEN.isoformat(),
        "group_count": 1,
        "total_parsed": 3,
        "rolecount": 3,
        "rolecount_available": True,
        "rolecount_match": True,
        "sections": [
            {
                "section": "board",
                "label": "Kårstyrelse",
                "assignments": [
                    Assignment("leader", "Kårordförande", "Kårordförande", "42", "Alva Ek", "board")
                ],
            }
        ],
        "vacancies": [],
    }


def test_render_html_has_header_and_rows():
    html = render_html("fortroende.html", _sample_context())
    assert "Scoutkåren Finn" in html
    assert "Kårordförande" in html and "Alva Ek" in html and "42" in html
    assert "Sida " in html and "counter(pages)" in html  # §19 page numbers
    assert "Höst 2026" in html


def test_html_to_pdf_is_guarded():
    html = render_html("fortroende.html", _sample_context())
    if HAS_WEASYPRINT:
        assert html_to_pdf(html)[:4] == b"%PDF"
    else:
        with pytest.raises(PdfUnavailable):
            html_to_pdf(html)


# --- API ------------------------------------------------------------------


@pytest.fixture
def client():
    settings = Settings(mode=Mode.FIXTURE, database_url="sqlite://", cohort_year=None)
    return create_app(settings).test_client()


def test_api_fortroende(client):
    d = client.get("/api/fortroende").get_json()
    assert d["assignments"]
    assert d["assignments"][0]["role_name"] == "Kårordförande"  # styrelse first
    names = [a["role_name"] for a in d["assignments"]]
    assert "Nyckelansvarig" in names  # unfamiliar role passes through
    assert d["group_count"] == len(d["assignments"])
    # Fast by default: the rolecount cross-check (slow organisation/group) is not
    # fetched unless asked for, so the blade renders without blocking.
    assert d["reconciliation"]["available"] is False
    assert d["reconciliation"]["requested"] is False
    # On demand it is available and agrees (fixture synthesises rolecount = all roles).
    agg = client.get("/api/fortroende?rolecount=1").get_json()["reconciliation"]
    assert agg["available"] is True and agg["match"] is True


def test_api_fortroende_xlsx(client):
    r = client.get("/api/fortroende.xlsx")
    assert r.status_code == 200 and r.data[:2] == b"PK"


def test_api_fortroende_pdf(client):
    r = client.get("/api/fortroende.pdf")
    if HAS_WEASYPRINT:
        assert r.status_code == 200 and r.data[:4] == b"%PDF"
    else:
        assert r.status_code == 503 and "WeasyPrint" in r.get_json()["error"]

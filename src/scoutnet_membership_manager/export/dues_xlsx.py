"""
Unpaid dues (Medlemsavgifter) as an Excel workbook — a chase list for
avdelningsledare (§19).

Everyone outstanding for the current or the previous term, following §4's three
buckets: the outstanding bucket is chased; ``not_invoiced`` and ``paid`` are
excluded; an **unrecognised** code is included too but flagged in the Anmärkning
column ("okänd betalkod – granska"). ``paid_partial_credit`` (paid the wrong
amount) is included and distinguished the same way.

A **single flat sheet** with everyone, sorted by avdelning then surname (Swedish
collation), so the operator can sort and filter it themselves. Streamed to the
browser, never written to disk server-side (§9).
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from scoutnet_membership_manager.collation import sort_key
from scoutnet_membership_manager.scoutnet.models import Member, MemberList, PaymentBucket

_BOLD = Font(bold=True)
_NO_AVDELNING = "(ingen avdelning)"
_PARTIAL_CODE = "paid_partial_credit"


def _route(m: Member) -> str | None:
    """
    Whether a member belongs on the list (§4/§19): 'review' if any term carries
    an unrecognised code, 'chase' if outstanding in either term, otherwise None
    (paid / not_invoiced only — excluded).
    """
    buckets = {m.current_payment(), m.prev_payment()}
    if PaymentBucket.UNKNOWN in buckets:
        return "review"
    if PaymentBucket.OUTSTANDING in buckets:
        return "chase"
    return None


def _headers(prev_label: str, cur_label: str) -> list[str]:
    prev = prev_label or "föregående termin"
    cur = cur_label or "innevarande termin"
    return [
        "Medlemsnummer",
        "Efternamn",
        "Förnamn",
        "Avdelning",
        "Utestående termin(er)",
        f"Status {prev}",
        f"Kod {prev}",
        f"Status {cur}",
        f"Kod {cur}",
        f"Förfallodatum {prev}",
        f"Förfallodatum {cur}",
        "KID",
        "Anmärkning",
        "E-post (mor)",
        "E-post (far)",
        "Mobil (mor)",
        "Mobil (far)",
        "Vårdnadshavare (mor)",
        "Vårdnadshavare (far)",
        "Medlemmens e-post",
        "Medlemmens mobil",
    ]


def _note(m: Member, pile: str) -> str:
    parts: list[str] = []
    if _PARTIAL_CODE in (m.current_term_code, m.prev_term_code):
        parts.append("Felaktigt belopp – korrigering, inte påminnelse")
    if pile == "review":
        parts.append("Okänd betalkod – granska, chasa inte")
    return " · ".join(parts)


def _row(m: Member, pile: str) -> list[object]:
    outstanding_terms = [
        label
        for bucket, label in (
            (m.prev_payment(), m.prev_term_label),
            (m.current_payment(), m.current_term_label),
        )
        if bucket is PaymentBucket.OUTSTANDING and label
    ]
    return [
        m.member_no,
        m.last_name,
        m.first_name,
        m.unit or _NO_AVDELNING,
        ", ".join(outstanding_terms),
        m.prev_term_value or "",
        m.prev_term_code or "",
        m.current_term_value or "",
        m.current_term_code or "",
        m.prev_term_due.display(),
        m.current_term_due.display(),
        m.kid or "",
        _note(m, pile),
        # Contacts (§10 dad/mum split). Empty cells are expected and are
        # themselves a data-quality signal (§11) — left blank, not "-".
        m.emails.get("contact_email_mum", ""),
        m.emails.get("contact_email_dad", ""),
        m.phones.get("contact_mobile_mum", ""),
        m.phones.get("contact_mobile_dad", ""),
        m.guardian_names.get("contact_mothers_name", ""),
        m.guardian_names.get("contact_fathers_name", ""),
        m.emails.get("contact_email", ""),
        m.phones.get("contact_mobile_phone", ""),
    ]


def build_dues_xlsx(memberlist: MemberList) -> bytes:
    """Build the unpaid-dues workbook as a single flat sheet (§19)."""
    rows = [(m, pile) for m in memberlist.members if (pile := _route(m)) is not None]
    rows.sort(
        key=lambda mp: (
            mp[0].unit is None,  # no-avdelning members last
            sort_key(mp[0].unit or ""),
            sort_key(mp[0].last_name),
            sort_key(mp[0].first_name),
        )
    )
    headers = _headers(memberlist.prev_term_label or "", memberlist.current_term_label or "")

    wb = Workbook()
    ws = wb.active
    ws.title = "Medlemsavgifter"
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h).font = _BOLD
    ws.freeze_panes = "A2"
    for r, (m, pile) in enumerate(rows, start=2):
        for c, value in enumerate(_row(m, pile), start=1):
            ws.cell(row=r, column=c, value=value)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

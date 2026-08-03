"""Uppflyttning changelist as an Excel workbook (§7 Phase 1).

Designed for the person doing the data entry: one sheet per target avdelning so
they work through one destination at a time, member_no first (the Scoutnet search
key), a done column, freeze panes, and a cover sheet with counts and the config
version. Off-cohort members must be acknowledged before export (§17); the ack is
stamped on the cover sheet.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from karverktyg.collation import sorted_sv
from karverktyg.uppflyttning.models import MasterSet

_BOLD = Font(bold=True)


class ChangelistAckRequired(RuntimeError):
    """Off-cohort members exist and have not been acknowledged (§17)."""


def _sheet_title(name: str) -> str:
    # Excel sheet titles are <=31 chars and cannot contain : \ / ? * [ ]
    safe = "".join(c for c in name if c not in ":\\/?*[]")
    return safe[:31] or "Avdelning"


def _write_cover(
    ws: Worksheet,
    master: MasterSet,
    kar_name: str,
    term_label: str | None,
    generated_at: datetime,
    config_version: str,
    ack_by: str | None,
    ack_at: datetime | None,
) -> None:
    ws.title = "Översikt"
    rows: list[tuple[str, object]] = [
        ("Kår", kar_name),
        ("Termin", term_label or ""),
        ("Uppflyttningsår (N)", master.cohort_year),
        ("Genererad", generated_at.isoformat(timespec="seconds")),
        ("Konfigurationsversion", config_version),
        ("", ""),
        ("Klara flyttar", len(master.ready())),
        ("Väntar på måldelning", len(master.pending())),
        ("Utanför årskull", len(master.off_cohort())),
        ("Undantagna (ledare/vuxna)", len(master.excluded())),
    ]
    if master.off_cohort():
        rows += [
            ("", ""),
            ("Utanför-årskull bekräftad av", ack_by or ""),
            ("Bekräftad", ack_at.isoformat(timespec="seconds") if ack_at else ""),
            ("Antal bekräftade", len(master.off_cohort())),
        ]
    rows += [("", ""), ("Per måldelning", "")]
    for target in sorted_sv(master.by_target().keys()):
        rows.append((f"  {target}", len(master.by_target()[target])))

    for r, (label, value) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=label).font = _BOLD
        ws.cell(row=r, column=2, value=value)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 28


def _write_target_sheet(ws: Worksheet, entries) -> None:
    headers = ["Medlemsnummer", "Namn", "Från", "Till", "Klar"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h).font = _BOLD
    ws.freeze_panes = "A2"
    for r, e in enumerate(sorted_sv(entries, key=lambda x: x.member_name), start=2):
        ws.cell(row=r, column=1, value=e.member_no)
        ws.cell(row=r, column=2, value=e.member_name)
        ws.cell(row=r, column=3, value=e.source_avdelning)
        ws.cell(row=r, column=4, value=e.target_avdelning)
        ws.cell(row=r, column=5, value="")  # done column
    for col, width in (("A", 16), ("B", 28), ("C", 16), ("D", 16), ("E", 8)):
        ws.column_dimensions[col].width = width


def build_changelist(
    master: MasterSet,
    *,
    kar_name: str,
    term_label: str | None,
    generated_at: datetime,
    config_version: str,
    ack_by: str | None = None,
    ack_at: datetime | None = None,
) -> bytes:
    if master.off_cohort() and not ack_by:
        raise ChangelistAckRequired(
            f"{len(master.off_cohort())} off-cohort member(s) must be acknowledged "
            "before the changelist can be exported (§17)"
        )

    wb = Workbook()
    _write_cover(
        wb.active, master, kar_name, term_label, generated_at, config_version, ack_by, ack_at
    )
    for target in sorted_sv(master.by_target().keys()):
        ws = wb.create_sheet(title=_sheet_title(target))
        _write_target_sheet(ws, master.by_target()[target])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
